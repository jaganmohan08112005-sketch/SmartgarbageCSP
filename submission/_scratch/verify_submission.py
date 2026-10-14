#!/usr/bin/env python3
"""One gate for the five R24 deliverables.

Re-checks every artifact in `submission/` against the contract it is submitted
under and against the source it was built from, then exits non-zero if any
deliverable drifted:

  1. book/SmartGarbage_BOOK_R24.pdf              page geometry + numbering
  2. documentation/...Softcopy.docx              DOCX container + content integrity
  3. ppt/SmartGarbage_Presentation.pptx          PPTX container + deck integrity
  4. video/SmartGarbage_Demo_Video.mp4           container/stream facts
  5. code/SmartGarbage_Source_Code.zip           manifest + per-file hashes

Two kinds of failure are reported, both fatal:

  * CONTRACT  - the artifact no longer satisfies the R24 requirement it was
                reviewed against (page size, numbering, slide/table counts,
                stream facts, archive contents).
  * DRIFT     - the artifact no longer matches the source it is built from: a
                code file changed after the zip was written (byte-level hash),
                or a generator changed after the artifact was committed.

Usage (from the repo root, or anywhere - paths resolve from this file):

    venv/Scripts/python.exe submission/_scratch/verify_submission.py
    venv/Scripts/python.exe submission/_scratch/verify_submission.py --selftest

`--selftest` proves the gate is not vacuous: it doctors a throwaway copy of
each artifact and asserts that the matching check fails on it.

Exit status: 0 = all five pass, 1 = at least one deliverable drifted.
"""
from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SUB = ROOT / "submission"

BOOK = SUB / "book" / "SmartGarbage_BOOK_R24.pdf"
DOCX = SUB / "documentation" / "SmartGarbage_Documentation_Softcopy.docx"
PPTX = SUB / "ppt" / "SmartGarbage_Presentation.pptx"
VIDEO = SUB / "video" / "SmartGarbage_Demo_Video.mp4"
ZIP = SUB / "code" / "SmartGarbage_Source_Code.zip"
MANIFEST = SUB / "_scratch" / "code_list.txt"

# ── The accepted contract (R24 template + submission/README.md) ──────────────
# Update these only after reviewing an intentional regeneration.
EXPECTED_BOOK_PAGES = 89
LETTER = (612.0, 792.0)          # 8.5 x 11 in in PDF points
MARGIN_PT = 72.0                 # 1 in
EXPECTED_DOCX_TABLES = 19
EXPECTED_DOCX_FIGURES = 20
EXPECTED_SLIDES = 20
EXPECTED_VIDEO_SECONDS = 590.0   # 9:50
VIDEO_TOLERANCE_S = 5.0

# Front matter order shared by the BOOK, the softcopy and the R24 contents list.
FRONT_ORDER = (
    "CERTIFICATE", "DECLARATION", "ACKNOWLEDGEMENT", "ABSTRACT",
    "LIST OF ABBREVIATIONS", "LIST OF FIGURES", "LIST OF TABLES", "TABLE OF CONTENTS",
)
REGISTERS = ("24331A4441", "24331A4434", "24331A4446", "24331A4426")
FRESH_FACTS = ("October, 2026", "2026-2027")
# A regenerated document that still carries these is stale, not just different.
STALE_FACTS = ("October, 2025", "2025-2026", "2433144441", "2433144446")

# What each artifact is built from, for the drift check.
SOURCES = {
    BOOK: ("report/generate_pdf.py", "report/report_figures.py",
           "report/_toc_targets.json", "report/_inputs/doc_template.docx"),
    DOCX: ("report/generate_final_report.py", "report/report_figures.py",
           "report/_inputs/doc_template.docx"),
    PPTX: ("report/build_deck_v2.py", "report/report_figures.py",
           "report/_inputs/ppt_template.pptx"),
    VIDEO: ("report/video/_make_video_live.py", "report/video/DEMO_VIDEO_SCRIPT.md"),
}
FORBIDDEN_ZIP = ("venv/", "node_modules/", "submission/", "audit/out/",
                 "report/DOCUMENTATION_UPDATED.docx", "report/DOCUMENTATION_UPDATED.pdf",
                 "report/DOCUMENT_FINAL.pptx")
MAX_PROBLEMS_SHOWN = 10


class Check:
    """One deliverable's verdict: facts measured, warnings, fatal problems."""

    def __init__(self, label: str, path: Path):
        self.label, self.path = label, path
        self.notes: list[str] = []
        self.warns: list[str] = []
        self.problems: list[str] = []

    def note(self, msg: str) -> None:
        self.notes.append(msg)

    def warn(self, msg: str) -> None:
        self.warns.append(msg)

    def fail(self, msg: str) -> None:
        self.problems.append(msg)

    @property
    def ok(self) -> bool:
        return not self.problems


# ── git helpers (used for the drift check) ───────────────────────────────────
def git(*args: str) -> str | None:
    try:
        r = subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True)
    except OSError:
        return None
    return r.stdout.strip() if r.returncode == 0 else None


def last_commit_epoch(rel: str) -> int | None:
    out = git("log", "-1", "--format=%ct", "--", rel)
    return int(out) if out and out.isdigit() else None


def dirty_paths() -> set[str]:
    """Paths with uncommitted changes, from `git status --porcelain -z`.

    Parsed here rather than through `git()`, which strips its output: a porcelain
    entry is two status characters plus a path, and an unstaged modification's
    first character is a SPACE, so stripping the blob shifts the first path by
    one character and silently drops the first dirty file from the set.
    """
    try:
        proc = subprocess.run(
            ["git", "status", "--porcelain", "-z", "--untracked-files=no"],
            cwd=ROOT, capture_output=True, text=True)
    except OSError:
        return set()
    if proc.returncode != 0:
        return set()
    return {e[3:] for e in proc.stdout.split("\0") if len(e) > 3}


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def check_drift(c: Check, rel_path: str, sources) -> None:
    """Fail when a generator changed after the artifact was committed."""
    t_art = last_commit_epoch(rel_path)
    if t_art is None:
        c.warn(f"{rel_path} is not committed yet - drift cannot be checked against HEAD")
        return
    dirty = dirty_paths()
    for src in sources:
        t_src = last_commit_epoch(src)
        if t_src is not None and t_src > t_art:
            c.fail(f"drift: {src} was committed after this artifact - regenerate and re-copy it")
        elif src in dirty:
            c.warn(f"{src} has uncommitted edits - rebuild this artifact before submitting")


# ── 1. BOOK: page geometry and numbering ─────────────────────────────────────
def _last_line(text: str) -> str:
    lines = [l.strip() for l in text.strip().splitlines() if l.strip()]
    return lines[-1] if lines else ""


def numbering_problems(pages: list[str], abbr: int) -> tuple[list[str], int]:
    """The R24 numbering rule, as a pure function of the page texts.

    Front matter before LIST OF ABBREVIATIONS carries no page number; LoA prints
    1 and every later footer is re-based from it. Returns (problems, pages whose
    footer was readable).
    """
    problems = []
    for i, text in enumerate(pages[: abbr - 1]):
        last = _last_line(text)
        if last.isdigit():
            problems.append(f"front-matter page {i + 1} prints a page number ({last!r})")
    loa_tail = _last_line(pages[abbr - 1])
    if loa_tail != "1":
        problems.append(f"LIST OF ABBREVIATIONS footer is not 1 (ends {loa_tail!r})")
    counted = 0
    for i in range(abbr - 1, len(pages)):
        last = _last_line(pages[i])
        if last.isdigit():
            counted += 1
            expected = i + 1 - (abbr - 1)
            if int(last) != expected:
                problems.append(f"page {i + 1}: footer {last} != the re-based {expected}")
    return problems, counted


def check_book(c: Check, path: Path = BOOK) -> None:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(p.extract_text() or "") for p in reader.pages]
    n = len(pages)
    c.note(f"{n} pages")
    if n != EXPECTED_BOOK_PAGES:
        c.fail(f"page count {n} != expected {EXPECTED_BOOK_PAGES} "
               "(review the change, then update EXPECTED_BOOK_PAGES)")

    for i, page in enumerate(reader.pages):
        box = page.mediabox
        w, h = float(box.width), float(box.height)
        if abs(w - LETTER[0]) > 0.5 or abs(h - LETTER[1]) > 0.5:
            c.fail(f"page {i + 1} is {w:.1f}x{h:.1f}pt, not US Letter 612x792pt")

    xs = []

    def visitor(_text, _cm, tm, _font, _size):
        if _text.strip():
            xs.append(tm[4])

    for page in reader.pages:
        page.extract_text(visitor_text=visitor)
    if not xs:
        c.fail("no extractable text positions - cannot verify the 1-inch margins")
    else:
        lo, hi = min(xs), max(xs)
        c.note(f"text band x {lo:.1f}..{hi:.1f}pt (1-inch margins = "
               f"{MARGIN_PT:.0f}..{LETTER[0] - MARGIN_PT:.0f})")
        if lo < MARGIN_PT - 1:
            c.fail(f"text starts at x={lo:.1f}pt, inside the 1-inch left margin")
        if hi > LETTER[0] - MARGIN_PT + 1:
            c.fail(f"text runs to x={hi:.1f}pt, past the 1-inch right margin")

    def head_lines(text: str, k: int = 3) -> list[str]:
        return [l.strip().lower() for l in text.strip().splitlines() if l.strip()][:k]

    def locate(title: str) -> int | None:
        key = title.lower()
        for i, text in enumerate(pages):
            if any(l == key for l in head_lines(text)):
                return i + 1
        return None

    # Front matter order (Figures before Tables, per the R24 contents list).
    found = {t: locate(t) for t in FRONT_ORDER}
    missing = [t for t, p in found.items() if p is None]
    if missing:
        c.fail(f"front-matter sections not found: {', '.join(missing)}")
    else:
        seq = [(t, found[t]) for t in FRONT_ORDER]
        if any(seq[i][1] >= seq[i + 1][1] for i in range(len(seq) - 1)):
            c.fail("front matter is out of order: "
                   + " -> ".join(f"{t}@{p}" for t, p in seq))
        else:
            c.note("front matter order " + " -> ".join(f"{t}={p}" for t, p in seq))

    abbr = found.get("LIST OF ABBREVIATIONS")

    # Numbering is re-based so LIST OF ABBREVIATIONS = 1.
    if abbr:
        problems, counted = numbering_problems(pages, abbr)
        for prob in problems:
            c.fail(prob)
        if not problems:
            if counted < (n - abbr + 1) * 0.8:
                c.warn(f"only {counted} of {n - abbr + 1} numbered pages yielded a footer")
            else:
                c.note(f"numbering re-based correctly on {counted} pages (LoA = 1)")

    all_text = "\n".join(pages)
    for stale in STALE_FACTS:
        if stale in all_text:
            c.fail(f"stale fact still printed in the book: {stale!r}")
    for fresh in FRESH_FACTS:
        if fresh not in all_text:
            c.fail(f"missing current fact: {fresh!r}")
    for reg in REGISTERS:
        if reg not in all_text:
            c.fail(f"register number missing from the book: {reg}")


# ── 2/3. DOCX and PPTX: container and content integrity ──────────────────────
def _assert_ooxml_container(c: Check, path: Path, parts) -> zipfile.ZipFile:
    if not zipfile.is_zipfile(path):
        raise ValueError("not a zip/OOXML container")
    zf = zipfile.ZipFile(path)
    corrupt = zf.testzip()
    if corrupt:
        c.fail(f"corrupt OOXML entry (CRC mismatch): {corrupt}")
    for part in parts:
        if part not in zf.namelist():
            c.fail(f"OOXML part missing: {part}")
            continue
        try:
            ET.fromstring(zf.read(part))
        except ET.ParseError as exc:
            c.fail(f"OOXML part does not parse: {part} ({exc})")
    return zf


def check_docx(c: Check, path: Path = DOCX) -> None:
    import docx

    zf = _assert_ooxml_container(
        c, path, ("[Content_Types].xml", "word/document.xml",
                  "word/_rels/document.xml.rels"))
    c.note(f"{len(zf.namelist())} OOXML parts")

    doc = docx.Document(str(path))
    tables, figures = len(doc.tables), len(doc.inline_shapes)
    c.note(f"{len(doc.paragraphs)} paragraphs, {tables} tables, {figures} figures")
    if tables != EXPECTED_DOCX_TABLES:
        c.fail(f"{tables} tables != expected {EXPECTED_DOCX_TABLES}")
    if figures != EXPECTED_DOCX_FIGURES:
        c.fail(f"{figures} figures != expected {EXPECTED_DOCX_FIGURES}")

    key = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    upper = [t.upper() for t in key]
    positions = {}
    for title in FRONT_ORDER:
        hits = [i for i, t in enumerate(upper) if t == title]
        if not hits:
            c.fail(f"front-matter section missing from the softcopy: {title}")
        else:
            positions[title] = hits[0]
    seq = [positions[t] for t in FRONT_ORDER if t in positions]
    if len(seq) == len(FRONT_ORDER) and any(a >= b for a, b in zip(seq, seq[1:])):
        c.fail("softcopy front matter is out of order: "
               + " -> ".join(f"{t}@{positions[t]}" for t in FRONT_ORDER))

    fonts = {}
    for para in doc.paragraphs:
        for run in para.runs:
            fonts[run.font.name] = fonts.get(run.font.name, 0) + 1
    named = {f: n for f, n in fonts.items() if f}
    if not named:
        c.fail("no explicit run fonts - the body must be Times New Roman")
    else:
        top, top_n = max(named.items(), key=lambda kv: kv[1])
        c.note(f"dominant font {top!r} on {top_n} of {sum(named.values())} named runs")
        if top != "Times New Roman":
            c.fail(f"dominant body font is {top!r}, not Times New Roman")

    text = "\n".join(key)
    for stale in STALE_FACTS:
        if stale in text:
            c.fail(f"stale fact still in the softcopy: {stale!r}")
    for fresh in FRESH_FACTS:
        if fresh not in text:
            c.fail(f"missing current fact in the softcopy: {fresh!r}")
    for reg in REGISTERS:
        if reg not in text:
            c.fail(f"register number missing from the softcopy: {reg}")


def check_pptx(c: Check, path: Path = PPTX) -> None:
    import pptx

    zf = _assert_ooxml_container(c, path, ("[Content_Types].xml", "ppt/presentation.xml"))
    c.note(f"{len(zf.namelist())} OOXML parts")

    prs = pptx.Presentation(str(path))
    slides = len(prs.slides)
    ratio = prs.slide_width / prs.slide_height
    c.note(f"{slides} slides, {prs.slide_width / 914400:.2f}x"
           f"{prs.slide_height / 914400:.2f} in ({ratio:.3f}:1)")
    if slides != EXPECTED_SLIDES:
        c.fail(f"{slides} slides != expected {EXPECTED_SLIDES}")
    if abs(ratio - 16 / 9) > 0.01:
        c.fail(f"slide aspect ratio {ratio:.3f} != 16:9")

    texts = []

    def add(frame) -> None:
        if frame.text.strip():
            texts.append(frame.text.strip())

    for slide in prs.slides:
        for shape in slide.shapes:
            # Register numbers and the student table live in a GraphicFrame, so
            # tables must be walked too - shape.text_frame is empty for them.
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        add(cell.text_frame)
            elif shape.has_text_frame:
                add(shape.text_frame)
    blob = "\n".join(texts)

    refs = [t for t in texts if re.match(r"\[\d+\]\s", t)]
    numbered = sorted(int(re.match(r"\[(\d+)\]", t).group(1)) for t in refs)
    if len(refs) < 10:
        c.fail(f"references slide carries only {len(refs)} numbered entries")
    elif numbered != list(range(1, len(numbered) + 1)):
        c.fail(f"IEEE reference numbering is not sequential: {numbered}")
    elif sum(1 for t in refs if re.search(r"\(\d{4}\)|et al\.", t)) < 3:
        c.fail("references are not IEEE style (no author/year entries)")
    else:
        c.note(f"{len(refs)} IEEE-numbered references, sequential from [1]")

    if not any(t.startswith("Thank You") for t in texts):
        c.fail("closing 'Thank You / Open Floor' slide is missing")
    for reg in REGISTERS:
        if reg not in blob:
            c.fail(f"register number missing from the deck: {reg}")
    for fresh in FRESH_FACTS:
        if fresh not in blob:
            c.warn(f"deck does not print {fresh!r}")


# ── 4. VIDEO: container and stream facts ─────────────────────────────────────
def check_video(c: Check, path: Path = VIDEO) -> None:
    import imageio_ffmpeg

    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    # -t 1 keeps the probe fast: the stream banner is printed before decoding.
    proc = subprocess.run(
        [ffmpeg, "-hide_banner", "-i", str(path), "-t", "1", "-f", "null", "-"],
        capture_output=True)
    blob = (proc.stdout + proc.stderr).decode("utf-8", "replace")

    dur = re.search(r"Duration: (\d+):(\d+):(\d+\.\d+)", blob)
    if not dur:
        c.fail("ffmpeg reports no readable duration - the container is unplayable")
        return
    seconds = int(dur.group(1)) * 3600 + int(dur.group(2)) * 60 + float(dur.group(3))
    c.note(f"{seconds:.2f}s ({int(seconds // 60)}:{seconds % 60:05.2f}), "
           f"{path.stat().st_size / 1e6:.1f} MB")
    if abs(seconds - EXPECTED_VIDEO_SECONDS) > VIDEO_TOLERANCE_S:
        c.fail(f"duration {seconds:.2f}s != {EXPECTED_VIDEO_SECONDS:.0f}s "
               f"(+/-{VIDEO_TOLERANCE_S:.0f}s)")

    video = re.search(r"Stream #\d+:\d+.*?: Video: (\w+).*?, (\d+)x(\d+)[,\s]", blob)
    audio = re.search(r"Stream #\d+:\d+.*?: Audio: (\w+).*?, (\d+) Hz, (\w+)", blob)
    if not video:
        c.fail("no video stream in the container")
    else:
        codec, w, h = video.group(1), int(video.group(2)), int(video.group(3))
        c.note(f"video {codec} {w}x{h}")
        if codec != "h264":
            c.fail(f"video codec is {codec}, not h264")
        if (w, h) != (1280, 720):
            c.fail(f"resolution {w}x{h}, not 1280x720")
    if not audio:
        c.fail("no audio stream in the container")
    else:
        c.note(f"audio {audio.group(1)} {audio.group(2)}Hz {audio.group(3)}")
        if audio.group(1) != "aac":
            c.fail(f"audio codec is {audio.group(1)}, not aac")


# ── 5. ZIP: manifest and per-file hashes ─────────────────────────────────────
def check_zip(c: Check, path: Path = ZIP, manifest: Path = MANIFEST) -> None:
    expected = [l.strip().replace("\r", "")
                for l in manifest.read_text(encoding="utf-8").splitlines() if l.strip()]
    zf = zipfile.ZipFile(path)
    names = zf.namelist()
    c.note(f"{len(names)} entries, {path.stat().st_size / 1e6:.1f} MB")

    dupes = sorted({n for n in names if names.count(n) > 1})
    if dupes:
        c.fail(f"archive carries duplicate entries: "
               + ", ".join(dupes[:MAX_PROBLEMS_SHOWN]))

    if names != expected:
        missing = sorted(set(expected) - set(names))
        extra = sorted(set(names) - set(expected))
        if missing:
            c.fail(f"archive is missing {len(missing)} manifest file(s): "
                   + ", ".join(missing[:MAX_PROBLEMS_SHOWN]))
        if extra:
            c.fail(f"archive carries {len(extra)} file(s) not in the manifest: "
                   + ", ".join(extra[:MAX_PROBLEMS_SHOWN]))
        if not missing and not extra:
            c.fail("archive order differs from the manifest order")

    forbidden = [n for n in names if any(f in n for f in FORBIDDEN_ZIP)]
    if forbidden:
        c.fail(f"archive carries paths that must never ship: "
               + ", ".join(forbidden[:MAX_PROBLEMS_SHOWN]))

    missing_on_disk, stale = [], []
    for name in dict.fromkeys(names):
        src = ROOT / name
        if not src.is_file():
            missing_on_disk.append(name)
        elif sha256(zf.read(name)) != sha256(src.read_bytes()):
            stale.append(name)
    if missing_on_disk:
        c.fail(f"{len(missing_on_disk)} archived file(s) no longer exist in the source "
               "tree: " + ", ".join(missing_on_disk[:MAX_PROBLEMS_SHOWN]))
    if stale:
        c.fail(f"archive drifted from the source tree - {len(stale)} file(s) differ; "
               "rebuild with build_code_zip.py --write: "
               + ", ".join(stale[:MAX_PROBLEMS_SHOWN]))
    if not stale and not missing_on_disk and names == expected:
        c.note(f"all {len(names)} entries hash-match the source tree")

    dirty = dirty_paths()
    pending = [n for n in names if n in dirty]
    if pending:
        c.warn(f"{len(pending)} archived file(s) have uncommitted edits (the archive "
               "matches the working tree, not HEAD) - commit before submitting")


# ── Runner ───────────────────────────────────────────────────────────────────
CHECKS = (
    ("BOOK PDF", "submission/book/SmartGarbage_BOOK_R24.pdf", check_book, BOOK),
    ("DOCX", "submission/documentation/SmartGarbage_Documentation_Softcopy.docx",
     check_docx, DOCX),
    ("PPTX", "submission/ppt/SmartGarbage_Presentation.pptx", check_pptx, PPTX),
    ("VIDEO", "submission/video/SmartGarbage_Demo_Video.mp4", check_video, VIDEO),
    ("ZIP", "submission/code/SmartGarbage_Source_Code.zip", check_zip, ZIP),
)


def run_one(index: int, label: str, rel: str, fn, artifact: Path) -> Check:
    print(f"\n[{index}/{len(CHECKS)}] {rel}  ({label})")
    c = Check(label, artifact)
    if not artifact.is_file():
        c.fail(f"deliverable is missing: {rel}")
    else:
        try:
            fn(c)
            if artifact in SOURCES:
                check_drift(c, rel, SOURCES[artifact])
        except ImportError as exc:
            c.fail(f"cannot verify - dependency missing ({exc.name}); "
                   f"pip install {exc.name}")
        except Exception as exc:  # a verification must fail loudly, never crash
            c.fail(f"cannot verify - {type(exc).__name__}: {exc}")
    for note in c.notes:
        print(f"      . {note}")
    for warn in c.warns:
        print(f"      ! {warn}")
    for prob in c.problems:
        print(f"      X {prob}")
    print(f"      {'PASS' if c.ok else 'FAIL'}")
    return c


def main() -> int:
    print("=" * 72)
    print("  R24 SUBMISSION VERIFICATION - five deliverables")
    print("=" * 72)
    results = [run_one(i, label, rel, fn, art)
               for i, (label, rel, fn, art) in enumerate(CHECKS, 1)]
    failed = [c for c in results if not c.ok]
    print("\n" + "=" * 72)
    if failed:
        print(f"  R24 SUBMISSION VERIFICATION: FAILED "
              f"({len(failed)} of {len(results)} deliverables)")
        print("=" * 72)
        for c in failed:
            print(f"  {c.label:9s} {c.path.relative_to(ROOT)}")
            for prob in c.problems:
                print(f"      X {prob}")
        return 1
    print(f"  R24 SUBMISSION VERIFICATION: ALL {len(results)} DELIVERABLES PASSED")
    print("=" * 72)
    return 0


def selftest() -> int:
    """Doctor a disposable copy of each artifact and assert the check fires."""
    print("=" * 72)
    print("  SELFTEST - each check must FAIL on a doctored artifact")
    print("=" * 72)
    tmp = Path(tempfile.mkdtemp(prefix="sg_verify_selftest_"))
    failures = []

    def report(name: str, c: Check) -> None:
        print(f"  {name:38s} {'detected' if not c.ok else 'MISSED'}")
        if c.ok:
            failures.append(name)
        else:
            for prob in c.problems[:4]:
                print(f"        X {prob}")

    def expect_fail(name: str, fn) -> None:
        c = Check(name, tmp)
        try:
            fn(c)
        except Exception as exc:
            c.fail(f"{type(exc).__name__}: {exc}")
        report(name, c)

    try:
        import pypdf
        from pypdf import PdfReader
        from pypdf.generic import RectangleObject

        writer = pypdf.PdfWriter()
        for page in PdfReader(str(BOOK)).pages:
            writer.add_page(page)
        writer.pages[0].mediabox = RectangleObject((0, 0, 595, 842))  # A4, not Letter
        a4 = tmp / "a4.pdf"
        with open(a4, "wb") as fh:
            writer.write(fh)
        expect_fail("book: A4 page geometry", lambda c: check_book(c, a4))

        trunc_docx = tmp / "trunc.docx"
        trunc_docx.write_bytes(DOCX.read_bytes()[:20_000])
        expect_fail("docx: truncated container", lambda c: check_docx(c, trunc_docx))

        trunc_pptx = tmp / "trunc.pptx"
        trunc_pptx.write_bytes(PPTX.read_bytes()[:20_000])
        expect_fail("pptx: truncated container", lambda c: check_pptx(c, trunc_pptx))

        head = tmp / "headonly.mp4"
        head.write_bytes(VIDEO.read_bytes()[:400_000])
        expect_fail("video: stream header stripped", lambda c: check_video(c, head))

        short = tmp / "short.mp4"
        import imageio_ffmpeg
        subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-v", "error", "-i",
                        str(VIDEO), "-t", "30", "-c", "copy", str(short), "-y"],
                       check=False)
        expect_fail("video: cut shortened to 30s", lambda c: check_video(c, short))

        # Numbering is a pure function of the page texts, so it can be falsified
        # directly instead of hand-editing a PDF's baked-in footers.
        def synth(*, numbered_front=False, wrong_footer=False) -> list[str]:
            pages = [f"CERTIFICATE\nissued {i}" for i in range(1, 5)]  # pages 1-4
            if numbered_front:
                pages[0] = "CERTIFICATE\n1"
            pages.append("LIST OF ABBREVIATIONS\na\n1")               # page 5 = LoA
            for i in range(6, 13):                                    # pages 6-12
                pages.append(f"body\n{99 if wrong_footer and i == 6 else i - 4}")
            return pages

        def expect_numbering_fail(name: str, pages: list[str]) -> None:
            # Exactly how check_book consumes the rule: every problem is fatal.
            c = Check(name, tmp)
            problems, _ = numbering_problems(pages, 5)
            for prob in problems:
                c.fail(prob)
            report(name, c)

        expect_numbering_fail("book: front matter prints numbers",
                              synth(numbered_front=True))
        expect_numbering_fail("book: footer breaks re-basing",
                              synth(wrong_footer=True))

        doctored = tmp / "source.zip"
        with zipfile.ZipFile(ZIP) as src, \
                zipfile.ZipFile(doctored, "w", zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "app/__init__.py":
                    data = b"# tampered after the archive was written\n"
                dst.writestr(info, data)
            dst.writestr("stray_note.txt", b"not in the manifest\n")
        expect_fail("zip: stray entry + tampered file",
                    lambda c: check_zip(c, doctored, MANIFEST))

        # Drift: point a fake source at the newest committed file in the tree.
        newest, newest_t = None, 0
        for rel in git("ls-files").splitlines():
            t = last_commit_epoch(rel)
            if t and t > newest_t:
                newest, newest_t = rel, t
        c = Check("drift", BOOK)
        check_drift(c, "submission/book/SmartGarbage_BOOK_R24.pdf",
                    (newest,) if newest else ())
        print(f"  {'drift: source newer than artifact':38s} "
              f"{'detected' if not c.ok else 'MISSED'}  (via {newest})")
        if c.ok:
            failures.append("drift")
        else:
            for prob in c.problems:
                print(f"        X {prob}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print("-" * 72)
    if failures:
        print(f"  SELFTEST FAILED - these checks did not fire: {', '.join(failures)}")
        return 1
    print("  SELFTEST PASSED - every check fails on its doctored artifact")
    return 0


if __name__ == "__main__":
    sys.exit(selftest() if "--selftest" in sys.argv else main())
